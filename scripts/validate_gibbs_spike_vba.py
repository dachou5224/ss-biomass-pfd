#!/usr/bin/env python3
"""本地复现 GibbsSpike.bas 数值逻辑（无 Excel VBA 编译器时的回归测试）。"""

from __future__ import annotations

import math
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from simulator.gibbs_spike import run_case1_rgpox_spike_comparison, run_case1_spike_comparison

XLSX = ROOT / "export" / "Gibbs_Spike_Test.xlsx"
SHEET_PREFIXES = ("", "RGX_")
N_SPEC = 8
N_Z = 3
PENALTY = 1.0e20


def _read_workbook_vectors(sheet_prefix: str = ""):
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws_setup = wb[f"{sheet_prefix}Setup"]
    setup = {
        ws_setup.cell(r, 1).value: ws_setup.cell(r, 2).value
        for r in range(2, 12)
        if ws_setup.cell(r, 1).value
    }
    t_k = float(setup["T_K"])
    p_bar = float(setup["P_bar"])
    rgas = float(setup["R_J_molK"])
    p_ratio = float(setup["P_ratio_bar"])

    ws = wb[f"{sheet_prefix}n_particular"]
    n_p = np.array([float(ws.cell(i + 2, 2).value) for i in range(N_SPEC)])
    ws = wb[f"{sheet_prefix}mu0"]
    mu0 = np.array([float(ws.cell(i + 2, 2).value) for i in range(N_SPEC)])
    ws = wb[f"{sheet_prefix}Null_B"]
    b = np.zeros((N_SPEC, N_Z))
    r = 2
    for i in range(N_SPEC):
        for k in range(N_Z):
            b[i, k] = float(ws.cell(r, 3).value)
            r += 1
    wb.close()
    return t_k, p_bar, rgas, p_ratio, n_p, b, mu0


def flows_from_z(z: np.ndarray, n_p: np.ndarray, bmat: np.ndarray) -> np.ndarray:
    n = n_p + bmat @ z
    return np.maximum(n, 1e-14)


def gibbs_objective(
    z: np.ndarray,
    n_p: np.ndarray,
    bmat: np.ndarray,
    mu0: np.ndarray,
    t_k: float,
    p_bar: float,
    rgas: float,
    p_ratio: float,
) -> float:
    n = flows_from_z(z, n_p, bmat)
    n_tot = float(n.sum())
    if n_tot <= 1e-18:
        return PENALTY
    pr = max(p_bar / p_ratio, 1e-12)
    y = np.maximum(n / n_tot * pr, 1e-18)
    mu = mu0 + rgas * t_k * np.log(y)
    val = float(np.sum(n * mu))
    if not math.isfinite(val):
        return PENALTY
    return val


def project_z_vec(z: np.ndarray, n_p: np.ndarray, bmat: np.ndarray) -> np.ndarray:
    z = z.copy()
    for _ in range(12):
        for i in range(N_SPEC):
            n_i = n_p[i] + bmat[i] @ z
            if n_i < 0.0:
                denom = float(np.dot(bmat[i], bmat[i])) + 1e-30
                shift = (-n_i + 1e-14) / denom
                z = z + shift * bmat[i]
    return z


def order_simplex(simplex: np.ndarray, f: np.ndarray) -> None:
    """与修复后 VBA 一致：用 float 排序，避免 Long 溢出。"""
    idx = np.argsort(f)
    simplex[:] = simplex[idx]
    f[:] = f[idx]


def optimize_nelder_mead(
    z0: np.ndarray,
    n_p: np.ndarray,
    bmat: np.ndarray,
    mu0: np.ndarray,
    t_k: float,
    p_bar: float,
    rgas: float,
    p_ratio: float,
    max_iter: int = 400,
) -> tuple[np.ndarray, float, int]:
    alpha, gamma, rho, sigma = 1.0, 2.0, 0.5, 0.5
    tol = 1e-8
    simplex = np.tile(z0, (4, 1))
    for i in range(1, 4):
        simplex[i, i - 1] += 0.05
        simplex[i] = project_z_vec(simplex[i], n_p, bmat)
    f = np.array(
        [gibbs_objective(simplex[i], n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio) for i in range(4)]
    )

    for it in range(1, max_iter + 1):
        order_simplex(simplex, f)
        if (f[3] - f[0]) < tol * (1.0 + abs(f[0])):
            return simplex[0], f[0], it
        centroid = simplex[:3].mean(axis=0)
        xr = project_z_vec(centroid + alpha * (centroid - simplex[3]), n_p, bmat)
        fr = gibbs_objective(xr, n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio)
        if fr < f[0] and fr >= f[1]:
            simplex[3] = xr
            f[3] = fr
        elif fr < f[0]:
            xe = project_z_vec(centroid + gamma * (xr - centroid), n_p, bmat)
            fe = gibbs_objective(xe, n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio)
            if fe < fr:
                simplex[3] = xe
                f[3] = fe
            else:
                simplex[3] = xr
                f[3] = fr
        else:
            xc = project_z_vec(centroid + rho * (simplex[3] - centroid), n_p, bmat)
            fc = gibbs_objective(xc, n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio)
            if fc < f[3]:
                simplex[3] = xc
                f[3] = fc
            else:
                for i in range(1, 4):
                    simplex[i] = project_z_vec(simplex[0] + sigma * (simplex[i] - simplex[0]), n_p, bmat)
                    f[i] = gibbs_objective(simplex[i], n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio)
    order_simplex(simplex, f)
    return simplex[0], f[0], max_iter


def _validate_block(label: str, sheet_prefix: str, cmp_fn) -> float:
    cmp = cmp_fn()
    gold = cmp["gold"]
    t_k, p_bar, rgas, p_ratio, n_p, bmat, mu0 = _read_workbook_vectors(sheet_prefix)

    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws_z = wb[f"{sheet_prefix}z_solution"]
    z_seed = np.array(
        [float(ws_z.cell(2, 2).value), float(ws_z.cell(3, 2).value), float(ws_z.cell(4, 2).value)]
    )
    wb.close()
    z_opt, f_opt, iters = optimize_nelder_mead(z_seed, n_p, bmat, mu0, t_k, p_bar, rgas, p_ratio)
    n_vba = flows_from_z(z_opt, n_p, bmat)

    print(f"\n=== {label} ===")
    print(f"Nelder-Mead iters={iters} objective={f_opt:.6e}")
    max_rel = 0.0
    for i, sp in enumerate(cmp["spike_case"].species):
        g = gold.species_flow_mol_h[sp]
        v = float(n_vba[i])
        rel = abs(v - g) / max(abs(g), 1e-3)
        if abs(g) >= 1e-3:
            max_rel = max(max_rel, rel)
        print(f"  {sp}: gold={g:.4f} vba_sim={v:.4f} rel={rel:.4%}")
    print(f"max_rel_err vs scipy gold: {max_rel:.4%}")
    return max_rel


def main() -> None:
    if not XLSX.is_file():
        print(f"缺少 {XLSX}，请先运行 build_gibbs_spike_workbook.py")
        sys.exit(1)

    max_all = 0.0
    max_all = max(max_all, _validate_block("INCI", "", run_case1_spike_comparison))
    max_all = max(max_all, _validate_block("RGPOX", "RGX_", run_case1_rgpox_spike_comparison))

    if max_all > 0.065:
        sys.exit(2)
    print("\nOK: INCI + RGPOX VBA 算法本地复现通过（可导入 GibbsSpike.bas / GibbsSpikeRgpox.bas）")


if __name__ == "__main__":
    main()
