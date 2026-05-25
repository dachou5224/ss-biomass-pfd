#!/usr/bin/env python3
"""
无 Excel 进程的 WebService 联调 CLI（Agent / CI 友好）。

复现 export/js/WebServiceDemo.js 的流程：
  命名区域 → POST /v1/compute/simulate-lite → 写回 Output_KPI_Table / Output_WS_Log_Table

Microsoft Excel / WPS 没有可在 shell 中执行 Office JS 的官方 headless CLI；
本脚本是推荐的自动化替代方案（openpyxl + curl）。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASE = "https://simapi.nice-ai.dev"
DEFAULT_WORKBOOK = ROOT / "export" / "Biomass_PFD_Simulator.xlsx"
LOG_TABLE_NAME = "Output_WS_Log_Table"
LOG_MAX_ROWS = 12
CURL_UA = "ss-biomass-pfd-excel-ws-cli/1.0"


class HeadlessRunResult:
    def __init__(
        self,
        *,
        ok: bool,
        health_status: int,
        post_status: int,
        api_status: str,
        kpi_rows: List[dict],
        log_rows: List[List[str]],
        workbook_path: Path,
        case_id: str,
        feed_count: int,
        error: str = "",
    ) -> None:
        self.ok = ok
        self.health_status = health_status
        self.post_status = post_status
        self.api_status = api_status
        self.kpi_rows = kpi_rows
        self.log_rows = log_rows
        self.workbook_path = workbook_path
        self.case_id = case_id
        self.feed_count = feed_count
        self.error = error


def _load_api_key(explicit: str = "", wb=None) -> str:
    if explicit.strip():
        return explicit.strip()
    env = os.getenv("SIM_API_KEY", "").strip()
    if env:
        return env
    if wb is not None:
        try:
            key = _read_named_scalar(wb, "Input_API_Key")
            if key:
                return key
        except KeyError:
            pass
    try:
        out = subprocess.check_output(
            [
                "ssh",
                "-o", "BatchMode=yes",
                "-o", "ConnectTimeout=15",
                "nice-ai-LZ",
                "grep SIM_API_KEY /etc/default/ss-biomass-api",
            ],
            text=True,
            timeout=30,
            stderr=subprocess.DEVNULL,
        )
        return out.strip().split("=", 1)[1]
    except (subprocess.CalledProcessError, FileNotFoundError, IndexError, subprocess.TimeoutExpired):
        return ""


def _curl(method: str, url: str, *, api_key: str = "", body: dict | None = None) -> Tuple[int, str]:
    cmd = ["curl", "-sS", "-A", CURL_UA, "-w", "\n%{http_code}", "-X", method, url]
    if api_key:
        cmd += ["-H", f"X-API-Key: {api_key}"]
    if body is not None:
        cmd += ["-H", "Content-Type: application/json", "-d", json.dumps(body)]
    proc = subprocess.run(cmd, capture_output=True, text=True, check=False)
    if proc.returncode != 0:
        raise RuntimeError(f"curl 失败: {proc.stderr.strip() or proc.stdout}")
    out = proc.stdout
    if "\n" not in out:
        raise RuntimeError(f"curl 响应异常: {out[:300]}")
    text, code_s = out.rsplit("\n", 1)
    return int(code_s), text


def _parse_defined_ref(attr_text: str) -> Tuple[str, int, int, int, int]:
    m = re.match(
        r"^(?:'([^']+)'|([^!]+))!\$([A-Z]+)\$(\d+)(?::\$([A-Z]+)\$(\d+))?$",
        attr_text.replace(" ", ""),
    )
    if not m:
        raise ValueError(f"无法解析命名区域引用: {attr_text}")
    sheet = m.group(1) or m.group(2)
    c1, r1 = m.group(3), int(m.group(4))
    c2, r2 = (m.group(5) or c1), int(m.group(6) or r1)

    def col_idx(col: str) -> int:
        n = 0
        for ch in col:
            n = n * 26 + (ord(ch) - 64)
        return n

    return sheet, col_idx(c1), r1, col_idx(c2), r2


def _read_named_scalar(wb, name: str) -> str:
    dn = wb.defined_names.get(name)
    if dn is None:
        raise KeyError(f"命名区域不存在: {name}")
    attr = dn.attr_text if hasattr(dn, "attr_text") else dn
    if isinstance(attr, list):
        attr = attr[0].attr_text
    sheet, c1, r1, c2, r2 = _parse_defined_ref(attr)
    ws = wb[sheet]
    return str(ws.cell(r1, c1).value or "").strip()


def _read_named_table(wb, name: str) -> List[List[Any]]:
    dn = wb.defined_names.get(name)
    if dn is None:
        raise KeyError(f"命名区域不存在: {name}")
    attr = dn.attr_text if hasattr(dn, "attr_text") else dn
    if isinstance(attr, list):
        attr = attr[0].attr_text
    sheet, c1, r1, c2, r2 = _parse_defined_ref(attr)
    ws = wb[sheet]
    return [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


def _write_named_table(wb, name: str, rows: List[List[Any]]) -> None:
    dn = wb.defined_names.get(name)
    if dn is None:
        raise KeyError(f"命名区域不存在: {name}")
    attr = dn.attr_text if hasattr(dn, "attr_text") else dn
    if isinstance(attr, list):
        attr = attr[0].attr_text
    sheet, c1, r1, c2, r2 = _parse_defined_ref(attr)
    ws = wb[sheet]
    nrows = r2 - r1 + 1
    ncols = c2 - c1 + 1
    for i in range(nrows):
        for j in range(ncols):
            val = rows[i][j] if i < len(rows) and j < len(rows[i]) else None
            ws.cell(r1 + i, c1 + j, val)


def _num(v: Any) -> float:
    try:
        x = float(v)
        return x if x == x else 0.0
    except (TypeError, ValueError):
        return 0.0


def _str(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _now_local() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def build_payload_from_workbook(wb) -> dict:
    case_id = _read_named_scalar(wb, "Input_CaseID") or "Case-1"
    feed_rows = _read_named_table(wb, "Input_Feed_Table")
    chem_rows = _read_named_table(wb, "Input_Chem_Table")

    pfd_feeds: Dict[str, dict] = {}
    for row in feed_rows:
        stream_id = _str(row[0] if row else "")
        if not stream_id or stream_id.startswith("Stream"):
            continue
        pfd_feeds[stream_id] = {
            "mass_kg_h": _num(row[1] if len(row) > 1 else 0),
            "temp_c": _num(row[2] if len(row) > 2 else 0),
            "pressure_bar": _num(row[3] if len(row) > 3 else 0),
        }

    o2in = {"O2": 95.0, "N2": 1.75, "Ar": 3.25}
    for row in chem_rows:
        field = _str(row[0] if row else "")
        value = _num(row[1] if len(row) > 1 else 0)
        if field == "O2IN O2 mol%":
            o2in["O2"] = value
        if field == "O2IN N2 mol%":
            o2in["N2"] = value
        if field == "O2IN Ar mol%":
            o2in["Ar"] = value

    return {"case_id": case_id, "pfd_feeds": pfd_feeds, "o2in_composition": o2in}


def build_payload_fixed() -> dict:
    return {
        "case_id": "Case-1",
        "pfd_feeds": {
            "Biomass": {"mass_kg_h": 4100, "temp_c": 25, "pressure_bar": 1.0},
            "O2IN": {"mass_kg_h": 1400, "temp_c": 25, "pressure_bar": 1.0},
        },
        "o2in_composition": {"O2": 95.0, "N2": 1.75, "Ar": 3.25},
    }


def _log_rows_from_steps(steps: List[Tuple[str, str]]) -> List[List[str]]:
    rows = [[_now_local(), step, detail] for step, detail in steps]
    while len(rows) < LOG_MAX_ROWS:
        rows.append(["", "", ""])
    return rows[:LOG_MAX_ROWS]


def run_headless_workbook_e2e(
    *,
    workbook: Path | None = None,
    base_url: str = DEFAULT_BASE,
    api_key: str = "",
    write_back: bool = True,
) -> HeadlessRunResult:
    """完整 headless 流程：health → 读表 → POST → 写回 KPI + 日志。"""
    from openpyxl import load_workbook

    path = workbook or DEFAULT_WORKBOOK
    if not path.is_file():
        raise FileNotFoundError(f"工作簿不存在: {path}")

    base = base_url.rstrip("/")
    logs: List[Tuple[str, str]] = [("START", base)]
    wb = load_workbook(path)
    key = _load_api_key(api_key, wb=wb)

    try:
        h_status, h_text = _curl("GET", f"{base}/health")
        logs.append(("GET /health", str(h_status)))
        if h_status != 200:
            logs.append(("ERROR", h_text[:200]))
            return HeadlessRunResult(
                ok=False,
                health_status=h_status,
                post_status=0,
                api_status="",
                kpi_rows=[],
                log_rows=_log_rows_from_steps(logs),
                workbook_path=path,
                case_id="",
                feed_count=0,
                error=f"health HTTP {h_status}",
            )

        if not key:
            logs.append(("ERROR", "缺少 API Key"))
            return HeadlessRunResult(
                ok=False,
                health_status=h_status,
                post_status=0,
                api_status="",
                kpi_rows=[],
                log_rows=_log_rows_from_steps(logs),
                workbook_path=path,
                case_id="",
                feed_count=0,
                error="missing API key",
            )

        logs.append(("API Key", f"已读取（{len(key)} 字符）"))
        payload = build_payload_from_workbook(wb)
        logs.append(("读取输入", f"case={payload['case_id']} feeds={len(payload['pfd_feeds'])}"))
        logs.append(("POST", f"{base}/v1/compute/simulate-lite"))

        p_status, p_text = _curl(
            "POST",
            f"{base}/v1/compute/simulate-lite",
            api_key=key,
            body=payload,
        )
        logs.append(("HTTP", str(p_status)))
        if p_status != 200:
            logs.append(("ERROR", p_text[:220]))
            if write_back and LOG_TABLE_NAME in wb.defined_names:
                _write_named_table(wb, LOG_TABLE_NAME, _log_rows_from_steps(logs))
                wb.save(path)
            return HeadlessRunResult(
                ok=False,
                health_status=h_status,
                post_status=p_status,
                api_status="",
                kpi_rows=[],
                log_rows=_log_rows_from_steps(logs),
                workbook_path=path,
                case_id=payload["case_id"],
                feed_count=len(payload["pfd_feeds"]),
                error=f"POST HTTP {p_status}",
            )

        data = json.loads(p_text)
        api_status = str(data.get("status", ""))
        kpi_rows = list(data.get("kpi_rows") or [])
        logs.append(("响应", f"status={api_status} kpi={len(kpi_rows)}"))
        ok = api_status == "ok" and len(kpi_rows) >= 5

        if write_back:
            if "Output_KPI_Table" in wb.defined_names:
                _write_named_table(
                    wb,
                    "Output_KPI_Table",
                    [[r.get("metric"), r.get("value"), r.get("unit")] for r in kpi_rows],
                )
                logs.append(("已写回", "Output_KPI_Table"))
            if LOG_TABLE_NAME in wb.defined_names:
                logs.append(("DONE" if ok else "WARN", "headless 完成"))
                _write_named_table(wb, LOG_TABLE_NAME, _log_rows_from_steps(logs))
            wb.save(path)

        return HeadlessRunResult(
            ok=ok,
            health_status=h_status,
            post_status=p_status,
            api_status=api_status,
            kpi_rows=kpi_rows,
            log_rows=_log_rows_from_steps(logs),
            workbook_path=path,
            case_id=payload["case_id"],
            feed_count=len(payload["pfd_feeds"]),
        )
    except Exception as exc:
        logs.append(("ERROR", str(exc)[:220]))
        if write_back and LOG_TABLE_NAME in wb.defined_names:
            try:
                _write_named_table(wb, LOG_TABLE_NAME, _log_rows_from_steps(logs))
                wb.save(path)
            except Exception:
                pass
        return HeadlessRunResult(
            ok=False,
            health_status=0,
            post_status=0,
            api_status="",
            kpi_rows=[],
            log_rows=_log_rows_from_steps(logs),
            workbook_path=path,
            case_id="",
            feed_count=0,
            error=str(exc),
        )


def read_kpi_metric(wb, metric: str) -> Optional[float]:
    rows = _read_named_table(wb, "Output_KPI_Table")
    for row in rows:
        if _str(row[0] if row else "") == metric:
            return _num(row[1] if len(row) > 1 else None)
    return None


def main() -> int:
    parser = argparse.ArgumentParser(description="Excel WebService 无头联调 CLI")
    parser.add_argument("--base-url", default=os.getenv("DEMO_BASE_URL", DEFAULT_BASE))
    parser.add_argument("--api-key", default="", help="或环境变量 SIM_API_KEY / 工作簿 Input_API_Key")
    parser.add_argument("--workbook", type=Path, default=None)
    parser.add_argument("--write-back", action="store_true", help="写回 KPI + 运行日志")
    parser.add_argument("--health-only", action="store_true")
    parser.add_argument("--e2e", action="store_true", help="完整 headless E2E（health + POST + 写回）")
    parser.add_argument("--fixed-payload", action="store_true")
    parser.add_argument("--json-out", type=Path, default=None)
    args = parser.parse_args()

    if args.e2e:
        result = run_headless_workbook_e2e(
            workbook=args.workbook,
            base_url=args.base_url,
            api_key=args.api_key,
            write_back=True,
        )
        print(f"headless E2E ok={result.ok} health={result.health_status} post={result.post_status}")
        print(f"case={result.case_id} feeds={result.feed_count} api_status={result.api_status}")
        for row in result.kpi_rows[:6]:
            print(f"  {row.get('metric')}\t{row.get('value')}\t{row.get('unit')}")
        if result.error:
            print(f"error: {result.error}", file=sys.stderr)
        return 0 if result.ok else 1

    base = args.base_url.rstrip("/")
    status, text = _curl("GET", f"{base}/health")
    print(f"GET /health -> {status}")
    print(text[:240])
    if status != 200:
        return 1
    if args.health_only:
        return 0

    from openpyxl import load_workbook

    wb = None
    api_key = _load_api_key(args.api_key)
    if args.workbook or (not args.fixed_payload and DEFAULT_WORKBOOK.is_file()):
        path = args.workbook or DEFAULT_WORKBOOK
        wb = load_workbook(path)
        api_key = _load_api_key(args.api_key, wb=wb)
        payload = build_payload_from_workbook(wb)
        print(f"已从工作簿读取: {path} case_id={payload['case_id']} feeds={len(payload['pfd_feeds'])}")
    else:
        payload = build_payload_fixed()
        print("使用内置固定 payload")

    if not api_key:
        print("错误: POST 需要 SIM_API_KEY", file=sys.stderr)
        return 1

    status, text = _curl("POST", f"{base}/v1/compute/simulate-lite", api_key=api_key, body=payload)
    print(f"POST /v1/compute/simulate-lite -> {status}")
    data = json.loads(text)
    print(f"status={data.get('status')} kpi_rows={len(data.get('kpi_rows', []))}")
    for row in data.get("kpi_rows") or []:
        print(f"  {row.get('metric')}\t{row.get('value')}\t{row.get('unit')}")

    if args.json_out:
        args.json_out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.write_back and wb is not None:
        kpi_rows = [[r.get("metric"), r.get("value"), r.get("unit")] for r in data.get("kpi_rows") or []]
        _write_named_table(wb, "Output_KPI_Table", kpi_rows)
        if LOG_TABLE_NAME in wb.defined_names:
            steps = [
                ("POST", str(status)),
                ("响应", str(data.get("status"))),
                ("DONE" if data.get("status") == "ok" else "WARN", "headless 写回"),
            ]
            _write_named_table(wb, LOG_TABLE_NAME, _log_rows_from_steps(steps))
        wb.save(args.workbook or DEFAULT_WORKBOOK)
        print(f"已写回 -> {args.workbook or DEFAULT_WORKBOOK}")

    return 0 if status == 200 and data.get("status") == "ok" else 1


if __name__ == "__main__":
    raise SystemExit(main())
