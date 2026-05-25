#!/usr/bin/env python3
"""生成 Gibbs 迁移 spike 工作簿：INCI + RGPOX，Python 金标准 vs 约化维数 (方案 B)。"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from simulator.gibbs_spike import run_case1_rgpox_spike_comparison, run_case1_spike_comparison
from simulator.parameters import EQUILIBRIUM_CFG, R_CONST

OUT_PATH = ROOT / "export" / "Gibbs_Spike_Test.xlsx"
VBA_INCI = ROOT / "export" / "vba" / "GibbsSpike.bas"
VBA_RGPOX = ROOT / "export" / "vba" / "GibbsSpikeRgpox.bas"


def _style_header(ws: Worksheet, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="334155")
    font = Font(bold=True, color="F8FAFC")
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")


def _write_df(ws: Worksheet, df: pd.DataFrame) -> None:
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    _style_header(ws)


def _append_spike_block(
    wb: Workbook,
    *,
    prefix: str,
    label: str,
    macro_name: str,
    cmp: Dict[str, Any],
    t_note: str,
) -> None:
    sc = cmp["spike_case"]
    gold = cmp["gold"]
    reduced = cmp["reduced"]
    p = prefix  # "" 或 "RGX_"

    ws_setup = wb.create_sheet(f"{p}Setup" if p else "Setup")
    setup = pd.DataFrame(
        [
            {"参数": "Case_ID", "值": sc.case_id, "说明": label},
            {"参数": "T_K", "值": sc.t_k, "说明": t_note},
            {"参数": "P_bar", "值": sc.p_bar, "说明": "系统压力 bar"},
            {"参数": "R_J_molK", "值": R_CONST, "说明": "气体常数"},
            {"参数": "P_ratio_bar", "值": EQUILIBRIUM_CFG["pressure_ratio_bar"], "说明": "化学势压力比"},
            {"参数": "balance_tol", "值": 1e-5, "说明": "元素守恒容差"},
            {"参数": "null_dim", "值": sc.null_basis.shape[1], "说明": "约化自由度"},
        ]
    )
    _write_df(ws_setup, setup)

    ws_b = wb.create_sheet(f"{p}Element_b")
    _write_df(ws_b, pd.DataFrame({"Element": list(sc.elements), "b_mol_h": sc.b_vector}))

    ws_a = wb.create_sheet(f"{p}Matrix_A")
    a_rows = [
        {"Element": el, "Species": sp, "A_ij": sc.a_matrix[i, j]}
        for i, el in enumerate(sc.elements)
        for j, sp in enumerate(sc.species)
    ]
    _write_df(ws_a, pd.DataFrame(a_rows))

    ws_np = wb.create_sheet(f"{p}n_particular")
    _write_df(ws_np, pd.DataFrame({"Species": list(sc.species), "n_p_mol_h": sc.n_particular}))

    ws_null = wb.create_sheet(f"{p}Null_B")
    null_rows = [
        {"Species": sp, "Basis": f"z{k + 1}", "B_ik": sc.null_basis[i, k]}
        for i, sp in enumerate(sc.species)
        for k in range(sc.null_basis.shape[1])
    ]
    _write_df(ws_null, pd.DataFrame(null_rows))

    ws_g0 = wb.create_sheet(f"{p}mu0")
    _write_df(ws_g0, pd.DataFrame({"Species": list(sc.species), "mu0_J_mol": sc.g0}))

    ws_cmp = wb.create_sheet(f"{p}Compare")
    rows = []
    for sp in sc.species:
        g = gold.species_flow_mol_h.get(sp, 0.0)
        r = reduced.species_flow_mol_h.get(sp, 0.0)
        rows.append(
            {
                "Species": sp,
                "Python_SLSQP": g,
                "Reduced_B_Python": r,
                "VBA_B": None,
                "RelErr_Reduced": abs(r - g) / max(abs(g), 1e-12),
                "RelErr_VBA": None,
            }
        )
    rows.append(
        {
            "Species": "_SUMMARY",
            "Python_SLSQP": gold.message,
            "Reduced_B_Python": reduced.message,
            "VBA_B": f"运行宏 {macro_name}",
            "RelErr_Reduced": cmp["max_rel_err"],
            "RelErr_VBA": None,
        }
    )
    rows.append(
        {
            "Species": "_balance_residual",
            "Python_SLSQP": "",
            "Reduced_B_Python": reduced.balance_residual,
            "VBA_B": None,
            "RelErr_Reduced": reduced.objective,
            "RelErr_VBA": gold.success,
        }
    )
    _write_df(ws_cmp, pd.DataFrame(rows))
    ws_cmp.column_dimensions["A"].width = 14
    ws_cmp.column_dimensions["B"].width = 18
    ws_cmp.column_dimensions["C"].width = 18
    ws_cmp.column_dimensions["D"].width = 18

    ws_z = wb.create_sheet(f"{p}z_solution")
    z_df = pd.DataFrame(
        [
            {"Variable": "z1", "Reduced_B": reduced.z_opt[0] if len(reduced.z_opt) > 0 else 0},
            {"Variable": "z2", "Reduced_B": reduced.z_opt[1] if len(reduced.z_opt) > 1 else 0},
            {"Variable": "z3", "Reduced_B": reduced.z_opt[2] if len(reduced.z_opt) > 2 else 0},
        ]
    )
    _write_df(ws_z, z_df)


def main() -> None:
    inci_cmp = run_case1_spike_comparison()
    rgpox_cmp = run_case1_rgpox_spike_comparison()

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    readme = [
        "Gibbs 迁移 Spike 工作簿（Case-1）",
        "",
        "【INCI @900°C】",
        "1. 表 Compare：Python_SLSQP vs Reduced_B_Python vs VBA_B",
        "2. 导入 export/vba/GibbsSpike.bas → 运行 RunGibbsSpikeCase1",
        f"   生成时 Reduced_B max rel err: {inci_cmp['max_rel_err']:.4%}",
        "",
        "【RGPOX @1400°C】全链元素进料 → 主物种 Gibbs",
        "1. 表 RGX_Compare：同上",
        "2. 导入 export/vba/GibbsSpikeRgpox.bas → 运行 RunGibbsSpikeRgpoxCase1",
        f"   生成时 Reduced_B max rel err: {rgpox_cmp['max_rel_err']:.4%}",
        "",
        "通过门槛建议: max_rel_err < 5%, balance_residual < 1e-4",
        "生成: python3 scripts/build_gibbs_spike_workbook.py",
    ]
    for i, line in enumerate(readme, start=1):
        ws_readme.cell(row=i, column=1, value=line)

    _append_spike_block(
        wb,
        prefix="",
        label="Case-1 INCI Gibbs 进料",
        macro_name="RunGibbsSpikeCase1",
        cmp=inci_cmp,
        t_note="INCI 温度 K",
    )
    _append_spike_block(
        wb,
        prefix="RGX_",
        label="Case-1 RGPOX Gibbs 元素进料（INCI+夹带+挥发+O2POX）",
        macro_name="RunGibbsSpikeRgpoxCase1",
        cmp=rgpox_cmp,
        t_note="RGPOX 固定 1400°C → K",
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"已写入 {OUT_PATH}")
    print(f"INCI  Reduced vs Gold: {inci_cmp['max_rel_err']:.4%}")
    print(f"RGPOX Reduced vs Gold: {rgpox_cmp['max_rel_err']:.4%}")
    print(f"INCI  VBA:  {VBA_INCI}")
    print(f"RGPOX VBA: {VBA_RGPOX}")


if __name__ == "__main__":
    main()
