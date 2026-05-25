"""Excel Spread Simulator 视觉主题与版式（流程模拟器风格）。"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# 色板（流程模拟器：深蓝标题 / 琥珀可编辑 / 灰只读输出）
CLR_TITLE_BG = "0F172A"
CLR_SECTION_BG = "1E40AF"
CLR_SECTION_FG = "FFFFFF"
CLR_TABLE_HEAD_BG = "334155"
CLR_TABLE_HEAD_FG = "F8FAFC"
CLR_EDIT_BG = "FFFBEB"
CLR_EDIT_BORDER = "D97706"
CLR_READONLY_BG = "F1F5F9"
CLR_READONLY_ALT = "E2E8F0"
CLR_KPI_BG = "ECFDF5"
CLR_KPI_BORDER = "059669"
CLR_CARD_BG = "FFFFFF"
CLR_CARD_BORDER = "94A3B8"
CLR_STREAM_HEAD = "B45309"
CLR_HINT_BG = "EFF6FF"
CLR_WARN_BG = "FEF2F2"
CLR_CONSOLE_BG = "1E293B"
CLR_CONSOLE_FG = "E2E8F0"
CLR_CONSOLE_HEAD = "334155"
CLR_CMD_BG = "F8FAFC"
CLR_STATUS_BG = "ECFDF5"
CLR_STATUS_BORDER = "059669"
CLR_HEALTH_GREEN_BG = "DCFCE7"
CLR_HEALTH_GREEN_FG = "166534"
CLR_HEALTH_YELLOW_BG = "FEF9C3"
CLR_HEALTH_YELLOW_FG = "854D0E"
CLR_HEALTH_RED_BG = "FEE2E2"
CLR_HEALTH_RED_FG = "991B1B"
CLR_HEALTH_GRAY_BG = "F1F5F9"
CLR_HEALTH_GRAY_FG = "64748B"

FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color=CLR_SECTION_FG)
FONT_TABLE_HEAD = Font(name="Calibri", size=10, bold=True, color=CLR_TABLE_HEAD_FG)
FONT_BODY = Font(name="Calibri", size=10)
FONT_KPI_LABEL = Font(name="Calibri", size=9, color="475569")
FONT_KPI_VALUE = Font(name="Calibri", size=14, bold=True, color="0F766E")
FONT_STREAM_ID = Font(name="Calibri", size=11, bold=True, color=CLR_STREAM_HEAD)
FONT_HINT = Font(name="Calibri", size=9, italic=True, color="64748B")
FONT_CONSOLE = Font(name="Consolas", size=9, color=CLR_CONSOLE_FG)
FONT_CONSOLE_HEAD = Font(name="Consolas", size=9, bold=True, color=CLR_CONSOLE_FG)
FONT_CMD = Font(name="Consolas", size=10, color="0F172A")
FONT_STATUS_VAL = Font(name="Calibri", size=11, bold=True, color="0F766E")

FILL_TITLE = PatternFill("solid", fgColor=CLR_TITLE_BG)
FILL_SECTION = PatternFill("solid", fgColor=CLR_SECTION_BG)
FILL_TABLE_HEAD = PatternFill("solid", fgColor=CLR_TABLE_HEAD_BG)
FILL_EDIT = PatternFill("solid", fgColor=CLR_EDIT_BG)
FILL_READONLY = PatternFill("solid", fgColor=CLR_READONLY_BG)
FILL_READONLY_ALT = PatternFill("solid", fgColor=CLR_READONLY_ALT)
FILL_KPI = PatternFill("solid", fgColor=CLR_KPI_BG)
FILL_CARD = PatternFill("solid", fgColor=CLR_CARD_BG)
FILL_STREAM_HEAD = PatternFill("solid", fgColor="FEF3C7")
FILL_HINT = PatternFill("solid", fgColor=CLR_HINT_BG)
FILL_WARN = PatternFill("solid", fgColor=CLR_WARN_BG)
FILL_CONSOLE = PatternFill("solid", fgColor=CLR_CONSOLE_BG)
FILL_CONSOLE_HEAD = PatternFill("solid", fgColor=CLR_CONSOLE_HEAD)
FILL_CMD = PatternFill("solid", fgColor=CLR_CMD_BG)
FILL_STATUS = PatternFill("solid", fgColor=CLR_STATUS_BG)
FILL_HEALTH_GREEN = PatternFill("solid", fgColor=CLR_HEALTH_GREEN_BG)
FILL_HEALTH_YELLOW = PatternFill("solid", fgColor=CLR_HEALTH_YELLOW_BG)
FILL_HEALTH_RED = PatternFill("solid", fgColor=CLR_HEALTH_RED_BG)
FILL_HEALTH_GRAY = PatternFill("solid", fgColor=CLR_HEALTH_GRAY_BG)

THIN = Side(style="thin", color="CBD5E1")
MEDIUM = Side(style="medium", color=CLR_EDIT_BORDER)
BORDER_TABLE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_EDIT = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
BORDER_CARD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

TAB_COLORS = {
    "Guide": "64748B",
    "WebService": "7C3AED",
    "PFD": "0D9488",
    "Model_Input": "2563EB",
    "Model_Output": "EA580C",
}

SPEC_UNITS: Dict[str, str] = {
    "INCI_T_C": "°C",
    "SLAG_T_C": "°C",
    "RGPOX_T_C": "°C",
    "SYSTEM_P_BAR": "bar",
}


@dataclass
class TableBlock:
    """工作表内一个数据区块（1-based 行号）。"""

    title: str
    header_row: int
    data_start: int
    data_end: int
    col_end: int = 6
    editable: bool = False
    value_cols: Tuple[int, ...] = (2,)
    number_format: Optional[str] = None


@dataclass
class SheetLayout:
    sheet: str
    title_row: int = 1
    col_end: int = 8
    hint_row: Optional[int] = None
    blocks: List[TableBlock] = field(default_factory=list)
    kpi_row_start: Optional[int] = None
    kpi_row_end: Optional[int] = None
    pfd_image_anchor: str = "A14"
    pfd_cards_origin: Tuple[int, int] = (4, 8)  # row, col for stream cards
    api_key_row: Optional[int] = None  # WebService 页 API Key 所在行（B 列）
    log_block: Optional["TableBlock"] = None  # WebService 页运行日志表
    health_block: Optional["TableBlock"] = None  # WebService API 健康监控表
    nav_row: Optional[int] = None  # 底部跨表导航行
    footer_row: Optional[int] = None  # 页脚提示行
    subtitle: Optional[str] = None  # 页眉副标题（format 时可选覆盖）


def _merge_row(ws: Worksheet, row: int, col_start: int, col_end: int, value: str, font: Font, fill: PatternFill) -> None:
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)
    cell = ws.cell(row=row, column=col_start, value=value)
    cell.font = font
    cell.fill = fill
    cell.alignment = ALIGN_LEFT if row > 1 else ALIGN_CENTER
    for c in range(col_start + 1, col_end + 1):
        ws.cell(row=row, column=c).fill = fill


def _set_row_heights(ws: Worksheet, row: int, height: float, col_end: int) -> None:
    ws.row_dimensions[row].height = height


def apply_page_title(ws: Worksheet, layout: SheetLayout, title: str, subtitle: str = "") -> None:
    col_end = max(layout.col_end, 6)
    layout.col_end = col_end
    _merge_row(ws, layout.title_row, 1, col_end, title, FONT_TITLE, FILL_TITLE)
    _set_row_heights(ws, layout.title_row, 32, col_end)
    if subtitle:
        r = layout.title_row + 1
        _merge_row(ws, r, 1, col_end, subtitle, FONT_HINT, FILL_HINT)
        _set_row_heights(ws, r, 18, col_end)


def apply_hint_bar(ws: Worksheet, row: int, col_end: int, text: str) -> None:
    _merge_row(ws, row, 1, col_end, text, FONT_HINT, FILL_HINT)
    _set_row_heights(ws, row, 22, col_end)


def apply_section_banner(ws: Worksheet, block: TableBlock) -> None:
    _merge_row(ws, block.header_row - 1, 1, block.col_end, block.title, FONT_SECTION, FILL_SECTION)
    _set_row_heights(ws, block.header_row - 1, 22, block.col_end)


def apply_table_block(ws: Worksheet, block: TableBlock, zebra: bool = True) -> None:
    apply_section_banner(ws, block)
    for c in range(1, block.col_end + 1):
        cell = ws.cell(row=block.header_row, column=c)
        if cell.value is not None:
            cell.font = FONT_TABLE_HEAD
            cell.fill = FILL_TABLE_HEAD
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_TABLE
    _set_row_heights(ws, block.header_row, 20, block.col_end)

    for r in range(block.data_start, block.data_end + 1):
        alt = zebra and (r - block.data_start) % 2 == 1
        base_fill = FILL_READONLY_ALT if alt and not block.editable else FILL_READONLY
        edit_fill = FILL_EDIT if block.editable else base_fill
        _set_row_heights(ws, r, 18, block.col_end)
        for c in range(1, block.col_end + 1):
            cell = ws.cell(row=r, column=c)
            if block.editable and c in block.value_cols:
                cell.fill = FILL_EDIT
                cell.border = BORDER_EDIT
                cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")
            else:
                cell.fill = edit_fill if block.editable else base_fill
                cell.border = BORDER_TABLE
                cell.font = FONT_BODY
            if block.number_format and c in block.value_cols and isinstance(cell.value, (int, float)):
                cell.number_format = block.number_format
            if c in block.value_cols and isinstance(cell.value, (int, float)):
                cell.alignment = ALIGN_RIGHT
            else:
                cell.alignment = ALIGN_LEFT


def apply_kpi_strip(ws: Worksheet, row_start: int, row_end: int, col_end: int = 8) -> None:
    """结果页顶部 KPI：指标 | 数值 | 单位 强化显示。"""
    for r in range(row_start, row_end + 1):
        lc = ws.cell(row=r, column=1)
        vc = ws.cell(row=r, column=2)
        uc = ws.cell(row=r, column=3)
        if lc.value:
            lc.font = FONT_KPI_LABEL
            lc.fill = FILL_KPI
            lc.alignment = ALIGN_LEFT
            lc.border = Border(
                left=Side(style="thin", color=CLR_KPI_BORDER),
                top=Side(style="thin", color=CLR_KPI_BORDER),
                bottom=Side(style="thin", color=CLR_KPI_BORDER),
            )
        if vc.value is not None and vc.value != "":
            vc.font = FONT_KPI_VALUE
            vc.fill = FILL_KPI
            vc.alignment = ALIGN_RIGHT
            vc.border = Border(
                top=Side(style="thin", color=CLR_KPI_BORDER),
                bottom=Side(style="thin", color=CLR_KPI_BORDER),
            )
            if isinstance(vc.value, (int, float)):
                label = str(lc.value or "")
                if "RMSD" in label or "误差" in label:
                    vc.number_format = "0.00"
                else:
                    vc.number_format = "#,##0.0"
        if uc.value:
            uc.font = FONT_KPI_LABEL
            uc.fill = FILL_KPI
            uc.alignment = ALIGN_CENTER
            uc.border = Border(
                right=Side(style="thin", color=CLR_KPI_BORDER),
                top=Side(style="thin", color=CLR_KPI_BORDER),
                bottom=Side(style="thin", color=CLR_KPI_BORDER),
            )
        _set_row_heights(ws, r, 28, col_end)


def write_stream_cards(
    ws: Worksheet,
    callouts: Sequence,
    start_row: int,
    start_col: int,
    cards_per_row: int = 2,
) -> Tuple[int, int]:
    """在 PFD 页写入流股卡片（模拟器物流标签）。返回 (末行, 末列)。"""
    card_w = 4
    card_h = 5
    row = start_row
    col = start_col
    idx = 0
    for c in callouts:
        if idx > 0 and idx % cards_per_row == 0:
            row += card_h + 1
            col = start_col
        base_row, base_col = row, col
        lines = c.as_rows()
        head = ws.cell(row=base_row, column=base_col, value=lines[0])
        head.font = FONT_STREAM_ID
        head.fill = FILL_STREAM_HEAD
        head.alignment = ALIGN_LEFT
        ws.merge_cells(
            start_row=base_row,
            start_column=base_col,
            end_row=base_row,
            end_column=base_col + card_w - 1,
        )
        for border_cell in range(base_col, base_col + card_w):
            ws.cell(row=base_row, column=border_cell).border = BORDER_CARD
        for i, text in enumerate(lines[1:], start=1):
            cell = ws.cell(row=base_row + i, column=base_col, value=text)
            cell.font = FONT_BODY
            cell.fill = FILL_CARD
            cell.alignment = ALIGN_LEFT
            ws.merge_cells(
                start_row=base_row + i,
                start_column=base_col,
                end_row=base_row + i,
                end_column=base_col + card_w - 1,
            )
            for cc in range(base_col, base_col + card_w):
                ws.cell(row=base_row + i, column=cc).fill = FILL_CARD
                ws.cell(row=base_row + i, column=cc).border = BORDER_CARD
        _set_row_heights(ws, base_row, 20, base_col + card_w)
        col += card_w + 1
        idx += 1
    end_row = row + card_h
    end_col = col
    return end_row, end_col


def apply_sheet_tab_color(ws: Worksheet, sheet_name: str) -> None:
    color = TAB_COLORS.get(sheet_name, "64748B")
    ws.sheet_properties.tabColor = color


def apply_column_widths(ws: Worksheet, widths: Dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def apply_view_options(ws: Worksheet, freeze: Optional[str] = None, zoom: int = 100) -> None:
    ws.sheet_view.zoomScale = zoom
    if freeze:
        ws.freeze_panes = freeze


def apply_hyperlink_nav(ws: Worksheet, links: Iterable[Tuple[str, str, int]]) -> None:
    """(显示文本, 目标 Sheet, 行)"""
    for text, target, row in links:
        cell = ws.cell(row=row, column=1, value=text)
        cell.hyperlink = f"#{target}!A1"
        cell.font = Font(name="Calibri", size=11, underline="single", color="2563EB")


def apply_workflow_block(ws: Worksheet, block: TableBlock) -> None:
    """Guide 流程表：步骤列强调。"""
    apply_table_block(ws, block, zebra=False)
    for r in range(block.data_start, block.data_end + 1):
        step = ws.cell(row=r, column=1)
        mod = ws.cell(row=r, column=2)
        if step.value:
            step.font = Font(name="Calibri", size=11, bold=True, color="1E40AF")
            step.alignment = ALIGN_CENTER
        if mod.value:
            mod.font = Font(name="Calibri", size=10, bold=True, color="0F172A")


def apply_bottom_nav_row(ws: Worksheet, layout: SheetLayout) -> None:
    if not layout.nav_row:
        return
    apply_hint_bar(ws, layout.nav_row, max(layout.col_end, 6), str(ws.cell(layout.nav_row, 1).value or ""))
    ws.merge_cells(
        start_row=layout.nav_row,
        start_column=1,
        end_row=layout.nav_row,
        end_column=min(layout.col_end, 6),
    )


def apply_guide_sheet(ws: Worksheet, layout: SheetLayout) -> None:
    apply_page_title(ws, layout, "生物质气化模拟器", "Spread Simulator · 工程首页")
    apply_hint_bar(
        ws,
        layout.hint_row or 3,
        layout.col_end,
        "标签顺序：Model_Input → WebService → Model_Output → PFD（与下方流程一致）",
    )
    apply_column_widths(ws, {"A": 8, "B": 16, "C": 28, "D": 22, "E": 12})
    for block in layout.blocks:
        if "流程" in (block.title or ""):
            apply_workflow_block(ws, block)
        else:
            apply_table_block(ws, block, zebra=False)
    apply_view_options(ws, freeze="A5", zoom=110)
    apply_sheet_tab_color(ws, "Guide")
    ws.sheet_view.showGridLines = False


def apply_input_sheet(ws: Worksheet, layout: SheetLayout) -> None:
    apply_page_title(ws, layout, "模型输入", layout.subtitle or "步骤 1/4")
    if layout.hint_row:
        apply_hint_bar(
            ws,
            layout.hint_row,
            layout.col_end,
            "黄色 = 可编辑输入 · 白色 = Excel 公式（便于审阅）",
        )
    apply_column_widths(
        ws,
        {"A": 30, "B": 16, "C": 12, "D": 28, "E": 16, "F": 14, "G": 16, "H": 16},
    )
    for block in layout.blocks:
        apply_table_block(ws, block)
    apply_bottom_nav_row(ws, layout)
    apply_view_options(ws, freeze="A6", zoom=100)
    apply_sheet_tab_color(ws, "Model_Input")
    ws.sheet_view.showGridLines = True


def apply_output_sheet(
    ws: Worksheet,
    layout: SheetLayout,
    default_widths: Dict[str, float],
) -> None:
    apply_page_title(ws, layout, "计算结果", layout.subtitle or "步骤 3/4 · 只读")
    if layout.hint_row:
        apply_hint_bar(ws, layout.hint_row, layout.col_end, "KPI 与组成表 · WebService 联调或完整仿真后更新")
    apply_column_widths(ws, default_widths)
    if layout.kpi_row_start and layout.kpi_row_end:
        apply_kpi_strip(ws, layout.kpi_row_start, layout.kpi_row_end, layout.col_end)
    for block in layout.blocks:
        fmt = "0.00" if "vol%" in (block.title or "") or "RMSD" in (block.title or "") else "#,##0.0"
        block.number_format = fmt if "vol%" in (block.title or "") else "#,##0.0"
        apply_table_block(ws, block, zebra=True)
    apply_bottom_nav_row(ws, layout)
    apply_view_options(ws, freeze="A6", zoom=100)
    apply_sheet_tab_color(ws, "Model_Output")
    ws.sheet_view.showGridLines = False


def apply_pfd_sheet(ws: Worksheet, layout: SheetLayout) -> None:
    apply_page_title(ws, layout, "工艺流程图", layout.subtitle or "步骤 4/4")
    if layout.hint_row:
        apply_hint_bar(
            ws,
            layout.hint_row,
            max(layout.col_end, 8),
            "左：流程示意图 · 右：流股卡片（随输入/结果刷新）",
        )
    apply_column_widths(
        ws,
        {"A": 2, "B": 14, "C": 36, "D": 14, "E": 14, "F": 14, "G": 20, "H": 20, "I": 20, "J": 20},
    )
    apply_view_options(ws, freeze="B5", zoom=90)
    apply_sheet_tab_color(ws, "PFD")
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def format_workbook(
    wb: Workbook,
    *,
    layouts: Dict[str, SheetLayout],
    guide_links_row: int = 8,
) -> None:
    """对工作簿各 Sheet 应用统一主题。"""
    default_widths = {"A": 22, "B": 16, "C": 14, "D": 14, "E": 16, "F": 14, "G": 18, "H": 18}

    if "Guide" in wb.sheetnames and "Guide" in layouts:
        apply_guide_sheet(wb["Guide"], layouts["Guide"])

    if "Model_Input" in wb.sheetnames and "Model_Input" in layouts:
        apply_input_sheet(wb["Model_Input"], layouts["Model_Input"])

    if "Model_Output" in wb.sheetnames and "Model_Output" in layouts:
        apply_output_sheet(wb["Model_Output"], layouts["Model_Output"], default_widths)

    if "PFD" in wb.sheetnames and "PFD" in layouts:
        apply_pfd_sheet(wb["PFD"], layouts["PFD"])

    if "WebService" in wb.sheetnames and "WebService" in layouts:
        apply_webservice_sheet(wb["WebService"], layouts["WebService"])


def apply_log_console(ws: Worksheet, block: TableBlock) -> None:
    """WebService 运行日志：终端风格深色底。"""
    apply_section_banner(ws, block)
    for c in range(1, block.col_end + 1):
        cell = ws.cell(row=block.header_row, column=c)
        if cell.value is not None:
            cell.font = FONT_CONSOLE_HEAD
            cell.fill = FILL_CONSOLE_HEAD
            cell.alignment = ALIGN_LEFT
    for r in range(block.data_start, block.data_end + 1):
        _set_row_heights(ws, r, 20, block.col_end)
        for c in range(1, block.col_end + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT_CONSOLE
            cell.fill = FILL_CONSOLE
            cell.alignment = ALIGN_LEFT if c <= 2 else ALIGN_LEFT
            cell.border = Border(
                left=Side(style="thin", color=CLR_CONSOLE_HEAD),
                right=Side(style="thin", color=CLR_CONSOLE_HEAD),
            )


def apply_command_block(ws: Worksheet, block: TableBlock) -> None:
    """运行命令区：等宽字体浅灰底。"""
    apply_section_banner(ws, block)
    for c in range(1, block.col_end + 1):
        cell = ws.cell(row=block.header_row, column=c)
        if cell.value is not None:
            cell.font = FONT_TABLE_HEAD
            cell.fill = FILL_TABLE_HEAD
    for r in range(block.data_start, block.data_end + 1):
        _set_row_heights(ws, r, 24, block.col_end)
        for c in range(1, block.col_end + 1):
            cell = ws.cell(row=r, column=c)
            if c == 1:
                cell.font = Font(name="Calibri", size=10, bold=True, color="475569")
                cell.fill = FILL_READONLY
            else:
                cell.font = FONT_CMD
                cell.fill = FILL_CMD
                cell.alignment = ALIGN_LEFT
            cell.border = BORDER_TABLE


def apply_health_monitor(ws: Worksheet, block: TableBlock) -> None:
    """API 健康监控：灯列 + 状态列（JS 运行时刷绿/黄/红）。"""
    apply_section_banner(ws, block)
    lamp_col = 2
    status_col = 3
    for r in range(block.data_start, block.data_end + 1):
        _set_row_heights(ws, r, 28, block.col_end)
        for c in range(1, block.col_end + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(
                left=Side(style="thin", color=CLR_STATUS_BORDER),
                right=Side(style="thin", color=CLR_STATUS_BORDER),
                top=Side(style="thin", color=CLR_STATUS_BORDER),
                bottom=Side(style="thin", color=CLR_STATUS_BORDER),
            )
            if c == 1:
                cell.font = Font(name="Calibri", size=10, bold=True, color="334155")
                cell.fill = FILL_READONLY
            elif c == lamp_col:
                cell.font = Font(name="Calibri", size=16, bold=True, color=CLR_HEALTH_GRAY_FG)
                cell.fill = FILL_HEALTH_GRAY
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c == status_col:
                cell.font = Font(name="Calibri", size=11, bold=True, color=CLR_HEALTH_GRAY_FG)
                cell.fill = FILL_HEALTH_GRAY
            else:
                cell.font = FONT_BODY
                cell.fill = FILL_READONLY
            cell.alignment = ALIGN_LEFT if c != lamp_col else Alignment(horizontal="center", vertical="center")


def apply_status_strip(ws: Worksheet, block: TableBlock) -> None:
    """顶部状态条：软件状态栏风格。"""
    apply_section_banner(ws, block)
    for r in range(block.data_start, block.data_end + 1):
        _set_row_heights(ws, r, 26, block.col_end)
        for c in range(1, block.col_end + 1):
            cell = ws.cell(row=r, column=c)
            if c in (1, 3, 5):
                cell.font = FONT_KPI_LABEL
                cell.fill = FILL_STATUS
            else:
                cell.font = FONT_STATUS_VAL
                cell.fill = FILL_STATUS
            cell.alignment = ALIGN_LEFT if c in (1, 3, 5) else ALIGN_LEFT
            cell.border = Border(
                left=Side(style="thin", color=CLR_STATUS_BORDER),
                right=Side(style="thin", color=CLR_STATUS_BORDER),
                top=Side(style="thin", color=CLR_STATUS_BORDER),
                bottom=Side(style="thin", color=CLR_STATUS_BORDER),
            )


def apply_api_key_row(ws: Worksheet, block: TableBlock) -> None:
    """API Key 行：加宽合并值单元格。"""
    apply_table_block(ws, block)
    for r in range(block.data_start, block.data_end + 1):
        label = str(ws.cell(row=r, column=1).value or "")
        if "API" in label or "密钥" in label:
            val_col = 2
            end_col = min(block.col_end, 6)
            if end_col > val_col:
                ws.merge_cells(
                    start_row=r,
                    start_column=val_col,
                    end_row=r,
                    end_column=end_col,
                )
            vc = ws.cell(row=r, column=val_col)
            vc.font = Font(name="Consolas", size=10, bold=True, color="92400E")
            vc.alignment = ALIGN_LEFT


def apply_webservice_sheet(ws: Worksheet, layout: SheetLayout) -> None:
    """WebService 计算客户端页：列宽、冻结、分区样式。"""
    apply_page_title(
        ws,
        layout,
        "在线计算服务",
        "Spread Simulator · Web 客户端 — 填密钥 → 运行命令 → 查看日志与结果",
    )
    if layout.hint_row:
        apply_hint_bar(
            ws,
            layout.hint_row,
            layout.col_end,
            "本页即计算软件「联调台」：无需打开 JS 编辑器，运行日志在下方自动刷新",
        )
    apply_column_widths(
        ws,
        {"A": 14, "B": 18, "C": 14, "D": 36, "E": 14, "F": 12, "G": 12, "H": 10},
    )
    for block in layout.blocks:
        title = block.title or ""
        if "运行日志" in title:
            apply_log_console(ws, block)
        elif "运行命令" in title:
            apply_command_block(ws, block)
        elif "健康监控" in title:
            apply_health_monitor(ws, block)
        elif "系统状态" in title:
            apply_status_strip(ws, block)
        elif "连接" in title or "授权" in title:
            apply_api_key_row(ws, block)
        else:
            apply_table_block(ws, block, zebra=False)
    if layout.nav_row:
        cell = ws.cell(row=layout.nav_row, column=1)
        cell.hyperlink = "#Model_Output!A1"
        cell.font = Font(name="Calibri", size=12, bold=True, underline="single", color="2563EB")
        cell.fill = FILL_HINT
        _set_row_heights(ws, layout.nav_row, 28, layout.col_end)
        ws.merge_cells(
            start_row=layout.nav_row,
            start_column=1,
            end_row=layout.nav_row,
            end_column=min(layout.col_end, 6),
        )
    if layout.footer_row:
        apply_hint_bar(
            ws,
            layout.footer_row,
            layout.col_end,
            "安全：API Key 等同密码 · 勿分享工作簿 · 勿提交 Git · 勿截图外传",
        )
    apply_view_options(ws, freeze="A6", zoom=105)
    apply_sheet_tab_color(ws, "WebService")
    ws.sheet_view.showGridLines = False
