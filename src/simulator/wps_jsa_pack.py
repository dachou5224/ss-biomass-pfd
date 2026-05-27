"""将 WPS JS 宏密封进 xlsx，生成可分发的 .xlsm（开箱即用）。"""

from __future__ import annotations

import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .parameters import PROJECT_ROOT

WPS_TEMPLATE_DIR = PROJECT_ROOT / "export" / "template" / "wps_jsa"
WPS_BOOTSTRAP_JS = WPS_TEMPLATE_DIR / "bootstrap.js"
WPS_PACKAGING_JSON = WPS_TEMPLATE_DIR / "packaging.json"
WPS_JS_SOURCE = PROJECT_ROOT / "export" / "js" / "WebServiceDemo.js"
MACRO_SOURCE_SHEET = "__MacroSrc__"

_EXPORTS_FOOTER = """
/* --- build: register public entrypoints for WPS bootstrap --- */
globalThis.__ss_exports__ = {
  validateWorkbookTemplate: validateWorkbookTemplate,
  refreshApiHealthMonitor: refreshApiHealthMonitor,
  runWebServiceHealthCheck: runWebServiceHealthCheck,
  runWebServiceLiteDemo: runWebServiceLiteDemo,
};
"""

MACRO_ENABLED_WORKBOOK_CT = (
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
)
STANDARD_WORKBOOK_CT = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


def load_packaging_config() -> dict[str, str]:
    defaults = {
        "jde_filename": "JDEData.bin",
        "jde_relationship_type": (
            "http://www.wps.cn/officeDocument/2020/relationships/jsProject"
        ),
        "jde_content_type": (
            "application/vnd.wps-officeDocument.spreadsheetml.jsProject"
        ),
    }
    if WPS_PACKAGING_JSON.is_file():
        data = json.loads(WPS_PACKAGING_JSON.read_text(encoding="utf-8"))
        defaults.update({k: str(v) for k, v in data.items() if not k.startswith("note")})
    return defaults


def build_macro_source_payload(js_path: Path | None = None) -> str:
    path = js_path or WPS_JS_SOURCE
    text = path.read_text(encoding="utf-8")
    if "__ss_exports__" not in text:
        text = text.rstrip() + _EXPORTS_FOOTER
    if len(text) > 32000:
        raise ValueError(
            f"联调脚本过长（{len(text)} 字符），超过单格上限；请拆分或精简 WebServiceDemo.js"
        )
    return text


def inject_macro_source_sheet(xlsx_bytes: bytes, js_path: Path | None = None) -> bytes:
    """在 xlsx 中写入隐藏表 __MacroSrc__（A1 存完整脚本）。"""
    payload = build_macro_source_payload(js_path)
    wb = load_workbook(BytesIO(xlsx_bytes))
    if MACRO_SOURCE_SHEET in wb.sheetnames:
        ws = wb[MACRO_SOURCE_SHEET]
        ws.delete_rows(1, ws.max_row or 1)
    else:
        ws = wb.create_sheet(MACRO_SOURCE_SHEET)
    ws["A1"] = payload
    ws.sheet_state = "veryHidden"
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def _next_rel_id(rels_xml: bytes) -> str:
    ids = [int(m.group(1)) for m in re.finditer(rb'Id="rId(\d+)"', rels_xml)]
    n = max(ids) + 1 if ids else 1
    return f"rId{n}"


def _patch_content_types(ct_xml: bytes, part_name: str, content_type: str) -> bytes:
    """字符串插入，避免 ElementTree 写出 ns0: 前缀导致 Excel/WPS 拒开。"""
    text = ct_xml.decode("utf-8")
    if part_name in text:
        return ct_xml
    override = (
        f'<Override PartName="{part_name}" '
        f'ContentType="{content_type}"/>'
    )
    if "</Types>" not in text:
        raise ValueError("Content_Types.xml 结构异常")
    return text.replace("</Types>", override + "</Types>", 1).encode("utf-8")


def _enable_macro_workbook_content_type(ct_xml: bytes) -> bytes:
    text = ct_xml.decode("utf-8")
    if MACRO_ENABLED_WORKBOOK_CT in text:
        return ct_xml
    if STANDARD_WORKBOOK_CT not in text:
        return ct_xml
    return text.replace(STANDARD_WORKBOOK_CT, MACRO_ENABLED_WORKBOOK_CT, 1).encode(
        "utf-8"
    )


def _patch_workbook_rels(rels_xml: bytes, target: str, rel_type: str) -> bytes:
    text = rels_xml.decode("utf-8")
    if f'Target="{target}"' in text:
        return rels_xml
    rid = _next_rel_id(rels_xml)
    rel = f'<Relationship Id="{rid}" Type="{rel_type}" Target="{target}"/>'
    if "</Relationships>" not in text:
        raise ValueError("workbook.xml.rels 结构异常")
    return text.replace("</Relationships>", rel + "</Relationships>", 1).encode("utf-8")


def seal_workbook_for_wps(
    xlsx_bytes: bytes,
    *,
    js_path: Path | None = None,
    bootstrap_path: Path | None = None,
) -> bytes:
    """
    把 bootstrap 写入 xl/JDEData.bin，并修补 OOXML，输出 WPS 可识别的 .xlsm 字节。
    """
    cfg = load_packaging_config()
    jde_name = cfg["jde_filename"]
    jde_target = jde_name
    part_name = f"/xl/{jde_name}"

    bootstrap = (bootstrap_path or WPS_BOOTSTRAP_JS).read_text(encoding="utf-8")
    sealed = inject_macro_source_sheet(xlsx_bytes, js_path=js_path)

    in_buf = BytesIO(sealed)
    out_buf = BytesIO()
    with zipfile.ZipFile(in_buf, "r") as zin, zipfile.ZipFile(
        out_buf, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = _enable_macro_workbook_content_type(data)
                data = _patch_content_types(
                    data, part_name, cfg["jde_content_type"]
                )
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = _patch_workbook_rels(
                    data, jde_target, cfg["jde_relationship_type"]
                )
            zout.writestr(
                item.filename,
                data,
                compress_type=item.compress_type,
            )
        zout.writestr(f"xl/{jde_name}", bootstrap.encode("utf-8"))
    return out_buf.getvalue()


def write_wps_simulator_workbook(
    xlsx_bytes: bytes,
    out_path: Path,
    *,
    js_path: Path | None = None,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_bytes(seal_workbook_for_wps(xlsx_bytes, js_path=js_path))
    return out_path
