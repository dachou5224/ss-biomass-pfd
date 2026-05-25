"""Excel WebService 无头 E2E（openpyxl + curl，无需打开 Excel）。"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "scripts" / "excel_ws_cli.py"
DEFAULT_WB = ROOT / "export" / "Biomass_PFD_Simulator.xlsx"
BUILD = [sys.executable, str(ROOT / "scripts" / "build_simulator_workbook.py"), "--case", "Case-1"]

requires_network = pytest.mark.skipif(
    os.getenv("SKIP_NETWORK_TESTS", "").lower() in ("1", "true", "yes"),
    reason="SKIP_NETWORK_TESTS 已设置",
)


def _fetch_api_key() -> str:
    key = os.getenv("SIM_API_KEY", "").strip()
    if key:
        return key
    try:
        out = subprocess.check_output(
            [
                "ssh", "-o", "BatchMode=yes", "-o", "ConnectTimeout=15",
                "nice-ai-LZ", "grep SIM_API_KEY /etc/default/ss-biomass-api",
            ],
            text=True,
            timeout=30,
            stderr=subprocess.DEVNULL,
        )
        return out.strip().split("=", 1)[1]
    except (subprocess.CalledProcessError, FileNotFoundError, IndexError, subprocess.TimeoutExpired):
        return ""


@pytest.fixture
def sim_api_key():
    key = os.getenv("SIM_API_KEY", "").strip()
    if key:
        return key
    key = _fetch_api_key()
    if not key:
        pytest.skip("无 SIM_API_KEY 且无法 SSH 读取 VPS")
    return key


def _ensure_workbook(tmp_path: Path) -> Path:
    if DEFAULT_WB.is_file():
        dest = tmp_path / "Biomass_PFD_Simulator.xlsx"
        shutil.copy2(DEFAULT_WB, dest)
        return dest
    subprocess.run(BUILD, cwd=ROOT, check=True)
    dest = tmp_path / "Biomass_PFD_Simulator.xlsx"
    shutil.copy2(DEFAULT_WB, dest)
    return dest


def test_workbook_has_webservice_named_ranges(tmp_path):
    wb_path = _ensure_workbook(tmp_path)
    wb = load_workbook(wb_path, read_only=True)
    for name in (
        "Input_CaseID",
        "Input_Feed_Table",
        "Input_Chem_Table",
        "Input_API_Key",
        "Output_KPI_Table",
        "Output_WS_Log_Table",
        "Output_API_Health_Table",
    ):
        assert name in wb.defined_names, f"缺少命名区域 {name}"


@requires_network
def test_excel_headless_e2e(tmp_path, sim_api_key, monkeypatch):
    """完整 headless：health → simulate-lite → 写回 KPI/日志。"""
    monkeypatch.setenv("SIM_API_KEY", sim_api_key)
    sys.path.insert(0, str(ROOT / "scripts"))
    from excel_ws_cli import read_kpi_metric, run_headless_workbook_e2e

    wb_path = _ensure_workbook(tmp_path)
    result = run_headless_workbook_e2e(workbook=wb_path, write_back=True)

    assert result.health_status == 200, result.error or result.log_rows
    assert result.post_status == 200, result.error or result.log_rows
    assert result.api_status == "ok"
    assert len(result.kpi_rows) >= 5

    wb = load_workbook(wb_path, read_only=True)
    total = read_kpi_metric(wb, "TOTAL_FEED_KG_H")
    assert total is not None and total > 0

    from excel_ws_cli import _read_named_table

    log_rows = _read_named_table(wb, "Output_WS_Log_Table")
    log_text = " ".join(str(c) for row in log_rows for c in row if c)
    assert "GET /health" in log_text or "START" in log_text


@requires_network
def test_cli_e2e_subprocess(sim_api_key, monkeypatch):
    monkeypatch.setenv("SIM_API_KEY", sim_api_key)
    if not DEFAULT_WB.is_file():
        subprocess.run(BUILD, cwd=ROOT, check=True)
    proc = subprocess.run(
        [sys.executable, str(CLI), "--e2e", "--workbook", str(DEFAULT_WB)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )
    assert proc.returncode == 0, proc.stderr + proc.stdout
